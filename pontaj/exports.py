"""Export XLSX lunar si saptamanal (doua foi: program + rezumat)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import grid as g

PALETTE = {
    'c-hrs': ('F7F7F5', '1F2328'), 'c-sp': ('E3F4E8', '1B6B3A'), 'c-sp05': ('DDF3EF', '0E6B5C'),
    'c-spp': ('FDEBD7', '9A4A00'), 'c-co': ('DFEBFB', '1D4F91'), 'c-con': ('FFF4CC', '7A5A00'),
    'c-cm': ('F8E1EE', '8E2A5E'), 'c-off': ('EEEEEC', '5F6368'), 'c-lp': ('DDF2F5', '0D6573'),
    'c-lfp': ('FBE3E1', '9B2C24'), 'c-in': ('F2F2F0', '9AA0A6'), 'c-empty': ('FFFFFF', 'BBBBBB'),
}
HEAD = ('E9EAEC', '3C4043')
TOTAL = ('DADCE0', '1F2328')
THIN = Side(style='thin', color='D0D3D8')
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')


def _style(cell, colors, bold=False, size=10, align=CENTER):
    bg, fg = colors
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.font = Font(name='Calibri', color=fg, bold=bold, size=size)
    cell.alignment = align
    cell.border = BORDER


def _section_colors(section):
    return section.bg.lstrip('#').upper(), section.color.lstrip('#').upper()


def _title(ws, text, width):
    ws.append([text])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    _style(ws.cell(1, 1), HEAD, bold=True, size=13, align=LEFT)
    ws.row_dimensions[1].height = 28
    ws.append([])


def _group_row(ws, label, width, colors=HEAD):
    ws.append([label])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    _style(ws.cell(r, 1), colors, bold=True, size=9, align=LEFT)


def _cell_text(cell, with_hours):
    if not cell.active:
        return 'IN'
    e = cell.entry
    if not e:
        return ''
    if e.is_interval:
        return f'{e.start_h:02d}-{e.end_h:02d}' + (f' ({e.hours}h)' if with_hours else '')
    if e.code == 'CO' and e.approved is not None:
        return 'CO (A)' if e.approved else 'CO (N)'
    return cell.label


def _program_sheet(ws, grid, title, headers, summary, with_hours):
    width = 2 + len(grid.days) + len(summary)
    _title(ws, title, width)
    ws.append(['Angajat', 'Sectie', *headers, *[s[0] for s in summary]])
    for i, c in enumerate(ws[ws.max_row], start=1):
        _style(c, HEAD, bold=True, size=9, align=LEFT if i <= 2 else CENTER)
    for shift, rows in grid.groups:
        _group_row(ws, f'Tura {shift}', width)
        for row in rows:
            values = [_cell_text(c, with_hours) for c in row.cells]
            ws.append([row.employee.name, row.employee.section.label, *values,
                       *[fn(row) for _, fn, _ in summary]])
            r = ws.max_row
            _style(ws.cell(r, 1), ('FFFFFF', '1F2328'), align=LEFT)
            _style(ws.cell(r, 2), _section_colors(row.employee.section), size=9)
            for i, c in enumerate(row.cells):
                _style(ws.cell(r, 3 + i), PALETTE[c.css], size=9)
            for j, (_, _, css) in enumerate(summary):
                _style(ws.cell(r, 3 + len(row.cells) + j), PALETTE.get(css, HEAD), bold=(j == 0))
    totals = _TotalsRow(grid)
    ws.append(['TOTAL', '', *[''] * len(grid.days), *[fn(totals) for _, fn, _ in summary]])
    r = ws.max_row
    for i in range(1, width + 1):
        _style(ws.cell(r, i), TOTAL, bold=True, align=LEFT if i <= 2 else CENTER)
    ws.freeze_panes = 'C4'


class _TotalsRow:
    """Adaptor ca functiile de rezumat sa mearga si pe randul de total."""

    def __init__(self, grid):
        self.stats = grid.totals
        self.bonus_label = ''


def _summary_sheet(ws, grid, title):
    cols = ['Angajat', 'Sectie', 'Total ore', 'SP', 'CO', 'CM', 'LFP']
    _title(ws, title, len(cols))
    ws.append(cols)
    for i, c in enumerate(ws[ws.max_row], start=1):
        _style(c, HEAD, bold=True, align=LEFT if i <= 2 else CENTER)
    for shift, rows in grid.groups:
        _group_row(ws, f'Tura {shift}', len(cols))
        for row in rows:
            s = row.stats
            ws.append([row.employee.name, row.employee.section.label, s.hours, s.sp, s.co, s.cm, s.lfp])
            r = ws.max_row
            _style(ws.cell(r, 1), ('FFFFFF', '1F2328'), align=LEFT)
            _style(ws.cell(r, 2), _section_colors(row.employee.section), size=9)
            _style(ws.cell(r, 3), PALETTE['c-hrs'], bold=True, size=11)
            for col, css in zip(range(4, 8), ('c-sp', 'c-co', 'c-cm', 'c-lfp')):
                _style(ws.cell(r, col), PALETTE[css])
    t = grid.totals
    ws.append(['TOTAL', '', t.hours, t.sp, t.co, t.cm, t.lfp])
    for i in range(1, len(cols) + 1):
        _style(ws.cell(ws.max_row, i), TOTAL, bold=True, align=LEFT if i <= 2 else CENTER)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 13
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 10
    ws.freeze_panes = 'C4'


def _num(v):
    return g.fmt_num(v)


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_xlsx(year, month):
    days = g.month_days(year, month)
    grid = g.build_grid(days, with_bonus=True)
    name = f'{g.MONTHS[month - 1]} {year}'
    wb = Workbook()
    ws = wb.active
    ws.title = f'Program {g.MONTHS[month - 1]}'
    summary = [
        ('Vara', lambda r: r.bonus_label, 'c-spp'),
        ('Ore', lambda r: r.stats.hours or '', 'c-hrs'),
        ('SP', lambda r: _num(r.stats.sp), 'c-sp'),
        ('CO', lambda r: r.stats.co or '', 'c-co'),
        ('CM', lambda r: r.stats.cm or '', 'c-cm'),
        ('LFP', lambda r: r.stats.lfp or '', 'c-lfp'),
    ]
    headers = [f'{g.DAY_SHORT[d.weekday()]}{d.day}' for d in days]
    _program_sheet(ws, grid, f'Program {name}', headers, summary, with_hours=True)
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 12
    for i in range(len(days) + len(summary)):
        ws.column_dimensions[ws.cell(3, 3 + i).column_letter].width = 7.5
    _summary_sheet(wb.create_sheet(f'Pontaj {g.MONTHS[month - 1]}'), grid, f'Pontaj {name}')
    return _save(wb), f'Pontaj_{g.MONTHS[month - 1]}_{year}.xlsx'


def week_xlsx(monday):
    days = g.week_days(monday)
    grid = g.build_grid(days)
    label = g.week_label(days)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Program saptamana'
    headers = [f'{g.DAY_LONG[d.weekday()]} {d.day}' for d in days]
    summary = [('Total ore', lambda r: r.stats.hours or 0, 'c-hrs')]
    _program_sheet(ws, grid, f'Program {label}', headers, summary, with_hours=False)
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 12
    for col in 'CDEFGHIJ':
        ws.column_dimensions[col].width = 12
    _summary_sheet(wb.create_sheet('Ore saptamana'), grid, f'Ore {label}')
    return _save(wb), f'Pontaj_Saptamana_{monday:%d_%m_%Y}.xlsx'
